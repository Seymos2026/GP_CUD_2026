"""
Bulk data import from an Excel workbook.

An admin downloads a template with three sheets (Faculty, Projects, Students),
fills it in, and uploads it. Rows are matched on email (people) and title
(projects): existing records are updated, new ones created.

New accounts get DEFAULT_PASSWORD and must_change_password=True, so the
middleware in accounts/middleware.py forces them to set their own password the
first time they sign in. Existing accounts keep the password they already have.

The whole import runs in one transaction: if any row is invalid, nothing is
saved and every problem is reported at once.
"""
from dataclasses import dataclass, field
from io import BytesIO

from django.contrib.auth import get_user_model
from django.db import transaction
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from accounts.models import Faculty, Student
from projects.models import FacultyProjectAssignment, Project

User = get_user_model()

DEFAULT_PASSWORD = "Temp@123"

SHEETS = {
    "Faculty": [
        ("Email", True, "m.injadat@cud.ac.ae"),
        ("First Name", True, "Mohammad"),
        ("Last Name", True, "Injadat"),
        ("Faculty ID", False, "FAC-001"),
        ("Department", False, "Computer Science"),
        ("Specialization", False, "Machine Learning"),
    ],
    "Projects": [
        ("Project Title", True, "GP_SP2025_01"),
        ("Description", False, "Smart campus navigation system"),
        ("Supervisor Email", True, "m.injadat@cud.ac.ae"),
        ("Judge Emails", False, "mehak.khurana@cud.ac.ae, najla.alfutaisi@cud.ac.ae"),
    ],
    "Students": [
        ("Email", True, "20220001588@students.cud.ac.ae"),
        ("First Name", True, "Lara"),
        ("Last Name", True, "Haddad"),
        ("Student ID", False, "20220001588"),
        ("Major", False, "Computer Science"),
        ("Project Title", False, "GP_SP2025_01"),
    ],
}

INSTRUCTIONS = [
    ("How to use this template", True),
    ("", False),
    ("1. Fill the Faculty sheet first, then Projects, then Students.", False),
    ("2. Columns marked * are required. Leave optional columns blank if unknown.", False),
    ("3. Delete the grey example row on each sheet before importing.", False),
    ("4. Upload the file on the Bulk Import page in the system.", False),
    ("", False),
    ("How rows are matched", True),
    ("", False),
    ("People are matched on Email, projects on Project Title.", False),
    ("If the record already exists it is updated. If not, it is created.", False),
    ("You can fix a mistake and re-import the same file safely.", False),
    ("", False),
    ("Passwords", True),
    ("", False),
    (f"New accounts are created with the password: {DEFAULT_PASSWORD}", False),
    ("The system forces each person to choose a new password at first login.", False),
    ("Existing accounts keep the password they already have.", False),
    ("", False),
    ("Supervisors and judges", True),
    ("", False),
    ("Supervisor Email must appear in the Faculty sheet or already exist in the system.", False),
    ("Judge Emails may list several people, separated by commas.", False),
    ("The same person cannot be both supervisor and judge on one project.", False),
]


@dataclass
class ImportResult:
    created: dict = field(default_factory=lambda: {"Faculty": 0, "Projects": 0, "Students": 0})
    updated: dict = field(default_factory=lambda: {"Faculty": 0, "Projects": 0, "Students": 0})
    errors: list = field(default_factory=list)

    def error(self, sheet, row, message):
        self.errors.append(f"{sheet} sheet, row {row}: {message}")

    @property
    def ok(self):
        return not self.errors

    @property
    def total_created(self):
        return sum(self.created.values())

    @property
    def total_updated(self):
        return sum(self.updated.values())


# ---------------------------------------------------------------------------
# Template
# ---------------------------------------------------------------------------

HEADER_FILL = PatternFill("solid", fgColor="4472C4")
REQUIRED_FILL = PatternFill("solid", fgColor="C55A11")
EXAMPLE_FILL = PatternFill("solid", fgColor="F2F2F2")


def build_template():
    """Return the import template as an in-memory .xlsx file."""
    wb = Workbook()

    ws = wb.active
    ws.title = "Instructions"
    for i, (text, is_heading) in enumerate(INSTRUCTIONS, start=1):
        cell = ws.cell(row=i, column=1, value=text)
        cell.font = Font(name="Arial", bold=is_heading, size=12 if is_heading else 11)
    ws.column_dimensions["A"].width = 90

    for sheet_name, columns in SHEETS.items():
        ws = wb.create_sheet(sheet_name)

        for col, (label, required, example) in enumerate(columns, start=1):
            cell = ws.cell(row=1, column=col, value=f"{label} *" if required else label)
            cell.font = Font(name="Arial", bold=True, color="FFFFFF")
            cell.fill = REQUIRED_FILL if required else HEADER_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center")

            example_cell = ws.cell(row=2, column=col, value=example)
            example_cell.font = Font(name="Arial", italic=True, color="808080")
            example_cell.fill = EXAMPLE_FILL

            ws.column_dimensions[get_column_letter(col)].width = max(len(label) + 6, 24)

        ws.freeze_panes = "A2"

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


# ---------------------------------------------------------------------------
# Import
# ---------------------------------------------------------------------------

def _rows(wb, sheet_name, columns, result):
    """Yield (row_number, {label: value}) for each non-empty, non-example row."""
    if sheet_name not in wb.sheetnames:
        result.errors.append(f"The workbook has no '{sheet_name}' sheet.")
        return

    ws = wb[sheet_name]
    headers = [str(c.value).replace("*", "").strip() if c.value else "" for c in ws[1]]
    expected = [label for label, _, _ in columns]

    missing = [label for label in expected if label not in headers]
    if missing:
        result.errors.append(
            f"{sheet_name} sheet is missing column(s): {', '.join(missing)}. "
            f"Download a fresh template."
        )
        return

    index = {label: headers.index(label) for label in expected}
    examples = {example for _, _, example in columns}

    for row_number, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        values = {}
        for label, position in index.items():
            value = row[position] if position < len(row) else None
            values[label] = str(value).strip() if value is not None else ""

        if not any(values.values()):
            continue
        # The untouched example row, left in by mistake.
        if all(values[label] in ("", example) for label, _, example in columns):
            continue

        yield row_number, values


def _require(values, labels, sheet, row_number, result):
    missing = [label for label in labels if not values.get(label)]
    if missing:
        result.error(sheet, row_number, f"missing required value(s): {', '.join(missing)}")
        return False
    return True


def _upsert_user(email, first, last, role, result):
    """Create or update the User behind a person row. Returns (user, created)."""
    user = User.objects.filter(username=email).first() or User.objects.filter(email=email).first()

    if user is None:
        user = User.objects.create(
            username=email, email=email, first_name=first, last_name=last,
            role=role, is_active=True, must_change_password=True,
        )
        user.set_password(DEFAULT_PASSWORD)
        user.save()
        return user, True

    user.email = email
    user.first_name = first
    user.last_name = last
    if user.role != role and not user.is_staff:
        user.role = role
    user.save()
    return user, False


def _import_faculty(wb, result):
    columns = SHEETS["Faculty"]
    for row_number, values in _rows(wb, "Faculty", columns, result):
        if not _require(values, ["Email", "First Name", "Last Name"], "Faculty", row_number, result):
            continue

        email = values["Email"].lower()
        user, created = _upsert_user(
            email, values["First Name"], values["Last Name"], User.Role.FACULTY, result,
        )

        defaults = {}
        if values["Faculty ID"]:
            defaults["faculty_id"] = values["Faculty ID"]
        if values["Department"]:
            defaults["department"] = values["Department"]
        if values["Specialization"]:
            defaults["specialization"] = values["Specialization"]

        profile = Faculty.objects.filter(user=user).first()
        if profile is None:
            Faculty.objects.create(user=user, **defaults)
        elif defaults:
            for key, value in defaults.items():
                setattr(profile, key, value)
            profile.save()

        result.created["Faculty" ] += 1 if created else 0
        result.updated["Faculty"] += 0 if created else 1


def _faculty_for(email, result, sheet, row_number, label):
    email = email.strip().lower()
    faculty = (
        Faculty.objects.filter(user__username__iexact=email).first()
        or Faculty.objects.filter(user__email__iexact=email).first()
    )
    if faculty is None:
        result.error(sheet, row_number, f"{label} '{email}' is not in the Faculty sheet or the system")
    return faculty


def _import_projects(wb, result, created_by):
    columns = SHEETS["Projects"]
    for row_number, values in _rows(wb, "Projects", columns, result):
        if not _require(values, ["Project Title", "Supervisor Email"], "Projects", row_number, result):
            continue

        title = values["Project Title"]
        project = Project.objects.filter(title__iexact=title).first()
        if project is None:
            project = Project.objects.create(
                title=title, description=values["Description"] or None, created_by=created_by,
            )
            result.created["Projects"] += 1
        else:
            if values["Description"]:
                project.description = values["Description"]
                project.save()
            result.updated["Projects"] += 1

        supervisor = _faculty_for(values["Supervisor Email"], result, "Projects", row_number, "Supervisor Email")

        judge_emails = [e for e in values["Judge Emails"].replace(";", ",").split(",") if e.strip()]
        judges = []
        for email in judge_emails:
            judge = _faculty_for(email, result, "Projects", row_number, "Judge email")
            if judge is not None:
                if supervisor is not None and judge.pk == supervisor.pk:
                    result.error(
                        "Projects", row_number,
                        f"'{email.strip()}' is listed as both supervisor and judge",
                    )
                else:
                    judges.append(judge)

        if supervisor is not None:
            FacultyProjectAssignment.objects.update_or_create(
                project=project, faculty=supervisor,
                defaults={"role": FacultyProjectAssignment.Role.SUPERVISOR},
            )
        for judge in judges:
            FacultyProjectAssignment.objects.update_or_create(
                project=project, faculty=judge,
                defaults={"role": FacultyProjectAssignment.Role.JUDGE},
            )


def _import_students(wb, result):
    columns = SHEETS["Students"]
    for row_number, values in _rows(wb, "Students", columns, result):
        if not _require(values, ["Email", "First Name", "Last Name"], "Students", row_number, result):
            continue

        email = values["Email"].lower()
        user, created = _upsert_user(
            email, values["First Name"], values["Last Name"], User.Role.STUDENT, result,
        )

        project = None
        if values["Project Title"]:
            project = Project.objects.filter(title__iexact=values["Project Title"]).first()
            if project is None:
                result.error(
                    "Students", row_number,
                    f"project '{values['Project Title']}' is not in the Projects sheet or the system",
                )
                continue

        profile = Student.objects.filter(user=user).first()
        if profile is None:
            profile = Student.objects.create(user=user)
        if values["Student ID"]:
            profile.student_id = values["Student ID"]
        if values["Major"]:
            profile.major = values["Major"]
        if project is not None:
            profile.project = project
        profile.save()

        result.created["Students"] += 1 if created else 0
        result.updated["Students"] += 0 if created else 1


def import_workbook(uploaded_file, created_by=None):
    """
    Read an uploaded .xlsx and apply it.

    Returns an ImportResult. If it carries any errors nothing was saved -
    the transaction is rolled back so a partial import can never happen.
    """
    result = ImportResult()

    try:
        wb = load_workbook(uploaded_file, data_only=True)
    except Exception as exc:
        result.errors.append(f"Could not read the file as an Excel workbook: {exc}")
        return result

    class _Rollback(Exception):
        pass

    try:
        with transaction.atomic():
            _import_faculty(wb, result)
            _import_projects(wb, result, created_by)
            _import_students(wb, result)
            if result.errors:
                raise _Rollback()
    except _Rollback:
        # Counts describe work that was rolled back; clear them to avoid confusion.
        result.created = {key: 0 for key in result.created}
        result.updated = {key: 0 for key in result.updated}

    return result
