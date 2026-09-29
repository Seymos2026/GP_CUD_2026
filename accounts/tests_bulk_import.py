from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import TestCase
from openpyxl import load_workbook

from accounts.bulk_import import DEFAULT_PASSWORD, SHEETS, build_template, import_workbook
from accounts.models import Faculty, Student
from projects.models import FacultyProjectAssignment, Project

User = get_user_model()


def make_workbook(faculty=(), projects=(), students=(), drop_example=True):
    """Fill a fresh template with the given rows and return it as an uploadable file."""
    wb = load_workbook(build_template())
    for sheet_name, rows in (("Faculty", faculty), ("Projects", projects), ("Students", students)):
        ws = wb[sheet_name]
        if drop_example:
            ws.delete_rows(2)
        for row in rows:
            ws.append(list(row))
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


class TemplateTests(TestCase):

    def test_template_has_the_expected_sheets_and_headers(self):
        wb = load_workbook(build_template())
        self.assertEqual(wb.sheetnames, ["Instructions", "Faculty", "Projects", "Students"])

        for sheet_name, columns in SHEETS.items():
            headers = [str(c.value).replace("*", "").strip() for c in wb[sheet_name][1]]
            self.assertEqual(headers, [label for label, _, _ in columns], sheet_name)

    def test_template_states_the_temporary_password(self):
        wb = load_workbook(build_template())
        text = "\n".join(str(row[0].value) for row in wb["Instructions"].iter_rows())
        self.assertIn(DEFAULT_PASSWORD, text)

    def test_an_untouched_template_imports_as_a_no_op(self):
        """The grey example rows must not create fake people."""
        wb = load_workbook(build_template())
        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        result = import_workbook(buffer)
        self.assertTrue(result.ok, result.errors)
        self.assertEqual(result.total_created, 0)
        self.assertEqual(User.objects.count(), 0)


class ImportTests(TestCase):

    FACULTY = [("sup@cud.ac.ae", "Mo", "Injadat", "FAC-001", "CS", "ML"),
               ("judge@cud.ac.ae", "Mehak", "Khurana", "FAC-002", "CS", "Security")]
    PROJECTS = [("GP_01", "Smart campus", "sup@cud.ac.ae", "judge@cud.ac.ae")]
    STUDENTS = [("s1@students.cud.ac.ae", "Lara", "Haddad", "20220001", "CS", "GP_01"),
                ("s2@students.cud.ac.ae", "Omar", "Adel", "20220002", "CS", "GP_01")]

    def test_full_round_trip(self):
        result = import_workbook(make_workbook(self.FACULTY, self.PROJECTS, self.STUDENTS))
        self.assertTrue(result.ok, result.errors)

        self.assertEqual(result.created, {"Faculty": 2, "Projects": 1, "Students": 2})
        self.assertEqual(Faculty.objects.count(), 2)
        self.assertEqual(Student.objects.count(), 2)

        project = Project.objects.get(title="GP_01")
        self.assertEqual(project.member_count, 2)
        self.assertEqual(project.supervisor.username, "sup@cud.ac.ae")
        self.assertEqual([j.username for j in project.judges], ["judge@cud.ac.ae"])

        student = Student.objects.get(user__username="s1@students.cud.ac.ae")
        self.assertEqual(student.student_id, "20220001")
        self.assertEqual(student.project, project)

    def test_new_users_get_the_temp_password_and_must_change_it(self):
        import_workbook(make_workbook(self.FACULTY, self.PROJECTS, self.STUDENTS))

        for username in ["sup@cud.ac.ae", "s1@students.cud.ac.ae"]:
            user = User.objects.get(username=username)
            self.assertTrue(user.check_password(DEFAULT_PASSWORD), username)
            self.assertTrue(user.must_change_password, username)

    def test_reimport_updates_and_never_resets_a_password(self):
        import_workbook(make_workbook(self.FACULTY, self.PROJECTS, self.STUDENTS))

        # The student has since chosen their own password.
        user = User.objects.get(username="s1@students.cud.ac.ae")
        user.set_password("TheirOwnPassword1!")
        user.must_change_password = False
        user.save()

        corrected = [("s1@students.cud.ac.ae", "Lara", "Haddad-Smith", "20220001", "CS", "GP_01")]
        result = import_workbook(make_workbook(self.FACULTY, self.PROJECTS, corrected))
        self.assertTrue(result.ok, result.errors)
        self.assertEqual(result.created["Students"], 0)
        self.assertEqual(result.updated["Students"], 1)

        user.refresh_from_db()
        self.assertEqual(user.last_name, "Haddad-Smith")
        self.assertTrue(user.check_password("TheirOwnPassword1!"))
        self.assertFalse(user.must_change_password)
        self.assertEqual(User.objects.count(), 4)

    def test_unknown_supervisor_rejects_the_whole_file(self):
        bad = [("GP_01", "Smart campus", "nobody@cud.ac.ae", "")]
        result = import_workbook(make_workbook(self.FACULTY, bad, self.STUDENTS))

        self.assertFalse(result.ok)
        self.assertIn("nobody@cud.ac.ae", result.errors[0])
        # Nothing at all was saved, including the valid Faculty rows.
        self.assertEqual(User.objects.count(), 0)
        self.assertEqual(Project.objects.count(), 0)

    def test_unknown_project_on_a_student_rejects_the_file(self):
        bad = [("s3@students.cud.ac.ae", "Sara", "Ali", "20220003", "CS", "GP_NOPE")]
        result = import_workbook(make_workbook(self.FACULTY, self.PROJECTS, bad))

        self.assertFalse(result.ok)
        self.assertIn("GP_NOPE", result.errors[0])
        self.assertEqual(User.objects.count(), 0)

    def test_missing_required_value_is_reported_with_its_row(self):
        bad = [("", "NoEmail", "Person", "", "", "")]
        result = import_workbook(make_workbook(bad, (), ()))

        self.assertFalse(result.ok)
        self.assertIn("Faculty sheet, row 2", result.errors[0])
        self.assertIn("Email", result.errors[0])

    def test_same_person_as_supervisor_and_judge_is_rejected(self):
        bad = [("GP_01", "Smart campus", "sup@cud.ac.ae", "sup@cud.ac.ae")]
        result = import_workbook(make_workbook(self.FACULTY, bad, ()))

        self.assertFalse(result.ok)
        self.assertIn("both supervisor and judge", result.errors[0])

    def test_a_non_excel_file_is_reported_not_crashed(self):
        result = import_workbook(BytesIO(b"this is not a spreadsheet"))
        self.assertFalse(result.ok)
        self.assertIn("Could not read the file", result.errors[0])


class ImportPageTests(TestCase):

    def test_a_superuser_gets_no_student_profile(self):
        """createsuperuser must not leave an admin looking like a student."""
        self.assertEqual(Student.objects.filter(user__username="admin").count(), 0)

    def setUp(self):
        admin = User.objects.create_superuser(
            username="admin", email="admin@cud.ac.ae", password="AdminPass123!",
        )
        User.objects.filter(pk=admin.pk).update(
            must_change_password=False, role=User.Role.ADMIN,
        )

    def test_admin_can_open_the_page_and_download_the_template(self):
        self.client.login(username="admin@cud.ac.ae", password="AdminPass123!")

        self.assertEqual(self.client.get("/bulk-import/").status_code, 200)

        response = self.client.get("/bulk-import/template/")
        self.assertEqual(response.status_code, 200)
        self.assertIn("GP_Import_Template.xlsx", response["Content-Disposition"])
        self.assertEqual(load_workbook(BytesIO(response.content)).sheetnames[0], "Instructions")

    def test_non_admin_is_turned_away(self):
        user = User.objects.create_user(
            username="stu@cud.ac.ae", email="stu@cud.ac.ae",
            password="StuPass123!", role=User.Role.STUDENT,
        )
        User.objects.filter(pk=user.pk).update(must_change_password=False)
        self.client.login(username="stu@cud.ac.ae", password="StuPass123!")

        self.assertEqual(self.client.get("/bulk-import/").status_code, 302)
        self.assertEqual(self.client.get("/bulk-import/template/").status_code, 302)

    def test_uploading_through_the_page_creates_records(self):
        self.client.login(username="admin@cud.ac.ae", password="AdminPass123!")

        upload = make_workbook(ImportTests.FACULTY, ImportTests.PROJECTS, ImportTests.STUDENTS)
        upload.name = "filled.xlsx"
        response = self.client.post("/bulk-import/", {"file": upload}, follow=True)

        self.assertEqual(response.status_code, 200)
        self.assertEqual(Project.objects.count(), 1)
        self.assertEqual(Student.objects.count(), 2)
        self.assertEqual(Faculty.objects.count(), 2)
