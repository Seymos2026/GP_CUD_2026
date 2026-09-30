from datetime import date
from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import TestCase
from openpyxl import load_workbook

from accounts.models import Faculty, Student
from projects.models import (
    FacultyProjectAssignment, Project, WeeklyAttendance, WeeklyProgress,
)

User = get_user_model()
PASSWORD = 'TestPass123!'


class WeeklyProgressTests(TestCase):
    """Weekly progress sheets: access control, attendance, and editing."""

    def setUp(self):
        self.project = Project.objects.create(title='Smart Campus')

        self.students = []
        for i in range(3):
            user = User.objects.create_user(
                username=f'stu{i}', email=f'stu{i}@cud.ac.ae',
                password=PASSWORD, role=User.Role.STUDENT,
            )
            User.objects.filter(pk=user.pk).update(must_change_password=False)
            student = user.student_profile
            student.project = self.project
            student.save()
            self.students.append(student)

        sup_user = User.objects.create_user(
            username='sup', email='sup@cud.ac.ae', password=PASSWORD,
            role=User.Role.FACULTY,
        )
        User.objects.filter(pk=sup_user.pk).update(must_change_password=False)
        self.supervisor = sup_user.faculty_profile
        FacultyProjectAssignment.objects.create(
            project=self.project, faculty=self.supervisor,
            role=FacultyProjectAssignment.Role.SUPERVISOR,
        )

        other_user = User.objects.create_user(
            username='judge', email='judge@cud.ac.ae', password=PASSWORD,
            role=User.Role.FACULTY,
        )
        User.objects.filter(pk=other_user.pk).update(must_change_password=False)
        self.judge = other_user.faculty_profile
        FacultyProjectAssignment.objects.create(
            project=self.project, faculty=self.judge,
            role=FacultyProjectAssignment.Role.JUDGE,
        )

        self.list_url = f'/projects/{self.project.id}/weekly/'
        self.create_url = f'/projects/{self.project.id}/weekly/new/'

    def login_supervisor(self):
        self.assertTrue(self.client.login(username='sup@cud.ac.ae', password=PASSWORD))

    # --- access control ----------------------------------------------------

    def test_supervisor_can_open_the_sheet(self):
        self.login_supervisor()
        self.assertEqual(self.client.get(self.list_url).status_code, 200)

    def test_judge_cannot_open_the_sheet(self):
        self.client.login(username='judge@cud.ac.ae', password=PASSWORD)
        response = self.client.get(self.list_url)
        self.assertRedirects(response, f'/projects/{self.project.id}/',
                             fetch_redirect_response=False)

    def test_student_cannot_open_the_sheet(self):
        self.client.login(username='stu0@cud.ac.ae', password=PASSWORD)
        response = self.client.get(self.list_url)
        self.assertRedirects(response, f'/projects/{self.project.id}/',
                             fetch_redirect_response=False)

    def test_admin_can_open_the_sheet(self):
        User.objects.create_superuser(
            username='admin', email='admin@cud.ac.ae', password=PASSWORD,
        )
        User.objects.filter(username='admin').update(
            must_change_password=False, role=User.Role.ADMIN,
        )
        self.client.login(username='admin@cud.ac.ae', password=PASSWORD)
        self.assertEqual(self.client.get(self.list_url).status_code, 200)

    # --- recording ---------------------------------------------------------

    def test_form_lists_every_assigned_student(self):
        self.login_supervisor()
        response = self.client.get(self.create_url)
        self.assertEqual(response.status_code, 200)
        for student in self.students:
            self.assertContains(response, f'attendance_{student.id}')

    def test_week_number_is_suggested(self):
        self.login_supervisor()
        self.assertEqual(self.client.get(self.create_url).context['form'].initial['week_number'], 1)

        WeeklyProgress.objects.create(
            project=self.project, week_number=4, meeting_date=date(2026, 9, 1),
        )
        self.assertEqual(self.client.get(self.create_url).context['form'].initial['week_number'], 5)

    def test_saving_records_attendance_and_notes(self):
        self.login_supervisor()
        response = self.client.post(self.create_url, {
            'week_number': 1,
            'meeting_date': '2026-09-29',
            'comments': 'Good progress on the API. Two members carried the work.',
            f'attendance_{self.students[0].id}': 'PRESENT',
            f'attendance_{self.students[1].id}': 'ABSENT',
            f'attendance_{self.students[2].id}': 'EXCUSED',
        })
        self.assertRedirects(response, self.list_url, fetch_redirect_response=False)

        report = WeeklyProgress.objects.get(project=self.project, week_number=1)
        self.assertEqual(report.meeting_date, date(2026, 9, 29))
        self.assertIn('Good progress', report.comments)
        self.assertEqual(report.supervisor, self.supervisor)
        self.assertEqual(report.attendance.count(), 3)
        self.assertEqual(report.present_count, 1)
        self.assertEqual(
            report.attendance.get(student=self.students[1]).status,
            WeeklyAttendance.Status.ABSENT,
        )

    def test_editing_updates_rather_than_duplicates(self):
        self.login_supervisor()
        self.client.post(self.create_url, {
            'week_number': 1, 'meeting_date': '2026-09-29', 'comments': 'First pass.',
            f'attendance_{self.students[0].id}': 'ABSENT',
        })
        report = WeeklyProgress.objects.get(project=self.project, week_number=1)

        self.client.post(f'/projects/{self.project.id}/weekly/{report.id}/', {
            'week_number': 1, 'meeting_date': '2026-09-29', 'comments': 'Corrected notes.',
            f'attendance_{self.students[0].id}': 'PRESENT',
        })

        report.refresh_from_db()
        self.assertEqual(WeeklyProgress.objects.filter(project=self.project).count(), 1)
        self.assertEqual(report.comments, 'Corrected notes.')
        self.assertEqual(report.attendance.count(), 3)
        self.assertEqual(
            report.attendance.get(student=self.students[0]).status,
            WeeklyAttendance.Status.PRESENT,
        )

    def test_duplicate_week_is_rejected(self):
        self.login_supervisor()
        WeeklyProgress.objects.create(
            project=self.project, week_number=2, meeting_date=date(2026, 9, 8),
        )
        response = self.client.post(self.create_url, {
            'week_number': 2, 'meeting_date': '2026-09-15', 'comments': 'Duplicate.',
        })
        self.assertEqual(response.status_code, 200)
        self.assertEqual(WeeklyProgress.objects.filter(project=self.project).count(), 1)

    def test_button_appears_for_supervisor_only(self):
        self.login_supervisor()
        response = self.client.get(f'/projects/{self.project.id}/')
        self.assertContains(response, 'Weekly Progress Sheet')

        self.client.login(username='judge@cud.ac.ae', password=PASSWORD)
        response = self.client.get(f'/projects/{self.project.id}/')
        self.assertNotContains(response, 'Weekly Progress Sheet')


class ProjectMembershipTests(TestCase):
    """Students attach straight to a Project now - no Team in between."""

    def setUp(self):
        self.project = Project.objects.create(title='Smart Campus')
        self.other = Project.objects.create(title='Unrelated')

    def _student(self, name, project):
        user = User.objects.create_user(
            username=name, email=f'{name}@cud.ac.ae', password=PASSWORD,
            role=User.Role.STUDENT,
        )
        User.objects.filter(pk=user.pk).update(must_change_password=False)
        student = user.student_profile
        student.project = project
        student.save()
        return student

    def test_members_returns_only_this_projects_students(self):
        a = self._student('a', self.project)
        b = self._student('b', self.project)
        self._student('c', self.other)

        self.assertEqual(set(self.project.members), {a, b})
        self.assertEqual(self.project.member_count, 2)

    def test_student_sees_only_their_own_project(self):
        self._student('a', self.project)
        self.client.login(username='a@cud.ac.ae', password=PASSWORD)

        response = self.client.get('/projects/')
        self.assertContains(response, 'Smart Campus')
        self.assertNotContains(response, 'Unrelated')

    def test_student_cannot_open_another_project(self):
        self._student('a', self.project)
        self.client.login(username='a@cud.ac.ae', password=PASSWORD)

        response = self.client.get(f'/projects/{self.other.id}/')
        self.assertRedirects(response, '/projects/', fetch_redirect_response=False)

    def test_unassigned_student_sees_nothing(self):
        user = User.objects.create_user(
            username='lost', email='lost@cud.ac.ae', password=PASSWORD,
            role=User.Role.STUDENT,
        )
        User.objects.filter(pk=user.pk).update(must_change_password=False)
        self.client.login(username='lost@cud.ac.ae', password=PASSWORD)

        response = self.client.get('/projects/')
        self.assertNotContains(response, 'Smart Campus')


class WeeklyProgressExportTests(TestCase):
    """The Excel export of a project's weekly sheets."""

    def setUp(self):
        self.project = Project.objects.create(title='Smart Campus')

        self.students = []
        for i in range(3):
            user = User.objects.create_user(
                username=f'stu{i}', email=f'stu{i}@cud.ac.ae',
                password=PASSWORD, role=User.Role.STUDENT,
                first_name=f'First{i}', last_name=f'Last{i}',
            )
            User.objects.filter(pk=user.pk).update(must_change_password=False)
            student = user.student_profile
            student.project = self.project
            student.student_id = f'2022000{i}'
            student.save()
            self.students.append(student)

        sup_user = User.objects.create_user(
            username='sup', email='sup@cud.ac.ae', password=PASSWORD,
            role=User.Role.FACULTY, first_name='Mo', last_name='Injadat',
        )
        User.objects.filter(pk=sup_user.pk).update(must_change_password=False)
        self.supervisor = sup_user.faculty_profile
        FacultyProjectAssignment.objects.create(
            project=self.project, faculty=self.supervisor,
            role=FacultyProjectAssignment.Role.SUPERVISOR,
        )

        # Two weeks recorded, with a different attendance pattern each week.
        for week, (meeting_date, statuses, note) in enumerate([
            (date(2026, 9, 8), ['PRESENT', 'PRESENT', 'ABSENT'], 'Kick-off. Scope agreed.'),
            (date(2026, 9, 15), ['PRESENT', 'EXCUSED', 'PRESENT'], 'API sketched out.'),
        ], start=1):
            report = WeeklyProgress.objects.create(
                project=self.project, supervisor=self.supervisor,
                week_number=week, meeting_date=meeting_date, comments=note,
            )
            for student, status in zip(self.students, statuses):
                WeeklyAttendance.objects.create(report=report, student=student, status=status)

        self.url = f'/projects/{self.project.id}/weekly/export/'

    def export(self):
        self.client.login(username='sup@cud.ac.ae', password=PASSWORD)
        response = self.client.get(self.url)
        self.assertEqual(response.status_code, 200)
        return response

    def test_downloads_as_an_xlsx_attachment(self):
        response = self.export()
        self.assertIn('spreadsheetml.sheet', response['Content-Type'])
        self.assertIn('attachment;', response['Content-Disposition'])
        self.assertIn('Smart_Campus_weekly_progress.xlsx', response['Content-Disposition'])

    def test_attendance_sheet_is_a_student_by_week_matrix(self):
        wb = load_workbook(BytesIO(self.export().content))
        self.assertEqual(wb.sheetnames, ['Attendance', 'Meeting Notes'])

        ws = wb['Attendance']
        headers = [c.value for c in ws[4]]
        self.assertEqual(headers, ['Student', 'Student ID', 'Week 1', 'Week 2', 'Present', 'Attendance %'])

        # Dates sit under the week headers.
        self.assertEqual(ws.cell(row=5, column=3).value.date(), date(2026, 9, 8))
        self.assertEqual(ws.cell(row=5, column=4).value.date(), date(2026, 9, 15))

        # Row 6 is the first student: present both weeks.
        self.assertEqual(ws.cell(row=6, column=1).value, 'First0 Last0')
        self.assertEqual(ws.cell(row=6, column=2).value, '20220000')
        self.assertEqual(ws.cell(row=6, column=3).value, 'Present')
        self.assertEqual(ws.cell(row=6, column=4).value, 'Present')

        # Third student: absent then present.
        self.assertEqual(ws.cell(row=8, column=3).value, 'Absent')
        self.assertEqual(ws.cell(row=8, column=4).value, 'Present')
        # Second student's excused week is recorded as such, not as absent.
        self.assertEqual(ws.cell(row=7, column=4).value, 'Excused')

    def test_totals_are_formulas_so_the_sheet_still_adds_up_if_edited(self):
        ws = load_workbook(BytesIO(self.export().content))['Attendance']

        present = ws.cell(row=6, column=5).value
        self.assertEqual(present, '=COUNTIF(C6:D6,"Present")')

        pct = ws.cell(row=6, column=6).value
        self.assertIn('COUNTIF(C6:D6,"Present")/COUNTA(C6:D6)', pct)
        self.assertEqual(ws.cell(row=6, column=6).number_format, '0%')

    def test_notes_sheet_carries_every_week(self):
        ws = load_workbook(BytesIO(self.export().content))['Meeting Notes']

        headers = [c.value for c in ws[4]]
        self.assertEqual(headers, ['Week', 'Date', 'Present', 'Total', 'Recorded By',
                                   'Progress and performance notes'])

        self.assertEqual(ws.cell(row=5, column=1).value, 1)
        self.assertEqual(ws.cell(row=5, column=3).value, 2)   # present count, week 1
        self.assertEqual(ws.cell(row=5, column=4).value, 3)   # total
        self.assertEqual(ws.cell(row=5, column=5).value, 'Mo Injadat')
        self.assertEqual(ws.cell(row=5, column=6).value, 'Kick-off. Scope agreed.')

        self.assertEqual(ws.cell(row=6, column=1).value, 2)
        self.assertEqual(ws.cell(row=6, column=3).value, 2)   # present + excused counted correctly
        self.assertEqual(ws.cell(row=6, column=6).value, 'API sketched out.')

    def test_judge_cannot_export(self):
        user = User.objects.create_user(
            username='judge', email='judge@cud.ac.ae', password=PASSWORD,
            role=User.Role.FACULTY,
        )
        User.objects.filter(pk=user.pk).update(must_change_password=False)
        FacultyProjectAssignment.objects.create(
            project=self.project, faculty=user.faculty_profile,
            role=FacultyProjectAssignment.Role.JUDGE,
        )
        self.client.login(username='judge@cud.ac.ae', password=PASSWORD)

        response = self.client.get(self.url)
        self.assertRedirects(response, f'/projects/{self.project.id}/',
                             fetch_redirect_response=False)

    def test_export_with_no_weeks_recorded_still_produces_a_file(self):
        WeeklyProgress.objects.all().delete()
        wb = load_workbook(BytesIO(self.export().content))

        self.assertEqual(wb['Attendance'].cell(row=4, column=3).value, 'Present')
        self.assertIn('No weekly sheets', str(wb['Meeting Notes'].cell(row=5, column=1).value))
